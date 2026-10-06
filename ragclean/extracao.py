"""Passos 1 e 2: conferir o corpus e extrair o texto dos documentos originais."""
from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path

from . import caminhos

FORMATOS = {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}
MIN_CARACTERES = 200  # abaixo disso o documento é tratado como "sem texto" (ex.: PDF escaneado)
CAMPOS = ["doc_id", "origem", "pasta", "sha256", "caracteres", "situacao", "texto"]


def sha256_arquivo(caminho: Path) -> str:
    h = hashlib.sha256()
    with caminho.open("rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def listar_corpus(corpus: Path | None = None) -> tuple[list[Path], list[Path]]:
    """Devolve (arquivos legíveis, arquivos ignorados) do corpus, em qualquer profundidade."""
    corpus = corpus or caminhos.CORPUS
    lidos, ignorados = [], []
    if corpus.exists():
        for arq in sorted(corpus.rglob("*")):
            if not arq.is_file() or arq.name.startswith("."):
                continue
            (lidos if arq.suffix.lower() in FORMATOS else ignorados).append(arq)
    return lidos, ignorados


def _pdf(caminho: Path) -> str:
    try:
        import pymupdf

        with pymupdf.open(caminho) as pdf:
            texto = "\n\n".join(pagina.get_text() for pagina in pdf)
        if len(texto.strip()) >= MIN_CARACTERES:
            return texto
    except Exception:
        pass
    try:
        import pdfplumber

        with pdfplumber.open(caminho) as pdf:
            return "\n\n".join(p.extract_text() or "" for p in pdf.pages)
    except Exception:
        return ""


def _docx(caminho: Path) -> str:
    import docx

    doc = docx.Document(str(caminho))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes.append(" | ".join(c.text for c in linha.cells))
    return "\n".join(partes)


def _texto_puro(caminho: Path) -> str:
    bruto = caminho.read_bytes()
    for codificacao in ("utf-8-sig", "latin-1"):
        try:
            return bruto.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return ""


def _html(caminho: Path) -> str:
    import html

    texto = _texto_puro(caminho)
    texto = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", texto)
    return html.unescape(re.sub(r"(?s)<[^>]+>", " ", texto))


def extrair_texto(caminho: Path) -> str:
    ext = caminho.suffix.lower()
    if ext == ".pdf":
        return _pdf(caminho)
    if ext == ".docx":
        return _docx(caminho)
    if ext in {".html", ".htm"}:
        return _html(caminho)
    return _texto_puro(caminho)


def ler_manifesto(manifesto: Path | None = None) -> list[dict[str, str]]:
    manifesto = manifesto or caminhos.MANIFESTO
    if not manifesto.exists():
        return []
    with manifesto.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def extrair_corpus(corpus: Path | None = None, avisar=print) -> list[dict[str, str]]:
    """Extrai o texto de cada documento e grava textos/ + manifesto.csv.

    Documentos cujo conteúdo não mudou (mesmo SHA-256) são reaproveitados.
    """
    corpus = corpus or caminhos.CORPUS
    caminhos.TEXTOS.mkdir(parents=True, exist_ok=True)
    anteriores = {r["sha256"]: r for r in ler_manifesto() if r["situacao"] != "duplicado"}
    arquivos, _ = listar_corpus(corpus)
    registros, vistos = [], set()

    for n, arq in enumerate(arquivos, start=1):
        origem = arq.relative_to(corpus).as_posix()
        sha = sha256_arquivo(arq)
        doc_id = "d" + sha[:10]
        partes = Path(origem).parts
        pasta = partes[0] if len(partes) > 1 else ""
        if sha in vistos:
            avisar(f"[{n}/{len(arquivos)}] {origem}: cópia idêntica de outro arquivo, ignorada")
            registros.append({"doc_id": doc_id, "origem": origem, "pasta": pasta, "sha256": sha,
                              "caracteres": "0", "situacao": "duplicado", "texto": ""})
            continue
        vistos.add(sha)
        destino = caminhos.TEXTOS / f"{doc_id}.txt"

        antigo = anteriores.get(sha)
        if antigo and destino.exists():
            registro = dict(antigo, origem=origem, pasta=pasta)
            avisar(f"[{n}/{len(arquivos)}] {origem}: já extraído, reaproveitado")
        else:
            try:
                texto = extrair_texto(arq)
                erro = ""
            except Exception as exc:
                texto, erro = "", f"erro: {type(exc).__name__}"
            texto = texto.replace("\x00", "")
            caracteres = len(texto.strip())
            situacao = erro or ("ok" if caracteres >= MIN_CARACTERES else "sem_texto")
            if situacao == "ok":
                destino.write_text(texto, encoding="utf-8")
            registro = {
                "doc_id": doc_id,
                "origem": origem,
                "pasta": pasta,
                "sha256": sha,
                "caracteres": str(caracteres),
                "situacao": situacao,
                "texto": destino.relative_to(caminhos.RAIZ).as_posix() if situacao == "ok" else "",
            }
            avisar(f"[{n}/{len(arquivos)}] {origem}: {situacao} ({caracteres:,} caracteres)".replace(",", "."))
        registros.append(registro)

    # remove textos de documentos que saíram do corpus
    validos = {r["doc_id"] for r in registros if r["situacao"] == "ok"}
    for antigo in caminhos.TEXTOS.glob("*.txt"):
        if antigo.stem not in validos:
            antigo.unlink()

    with caminhos.MANIFESTO.open("w", encoding="utf-8", newline="") as f:
        escritor = csv.DictWriter(f, fieldnames=CAMPOS)
        escritor.writeheader()
        escritor.writerows(registros)
    return registros
