from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook

from rag_project.question_catalog import get_questions
from rag_project.corpus_registry import load_registry


def build_workbook(input_dir: str | Path = 'analysis/countries', output_path: str | Path = '02_perguntas_respostas.xlsx') -> Path:
    root = Path(input_dir)
    rows = []
    for path in sorted(root.glob('*/respostas.csv')):
        with path.open(encoding='utf-8', newline='') as handle:
            rows.extend(csv.DictReader(handle))
    questions = {q['question_id']: q for q in get_questions()}
    countries = [c for c in load_registry().get('countries', []) if c.get('include_in_analysis')]
    workbook = Workbook()
    dictionary = workbook.active
    dictionary.title = '00_dicionario'
    dictionary.append(['campo', 'descrição'])
    for field, description in [('question_id', 'Código estável da pergunta'), ('country', 'País analisado'), ('response', 'Resposta e evidências'), ('validation_status', 'Estado da validação humana')]:
        dictionary.append([field, description])
    long_sheet = workbook.create_sheet('01_respostas_longas')
    fields = ['question_id', 'question_text', 'country_code', 'country', 'answer_summary', 'evidence_classification', 'source_ids', 'document_ids', 'pages', 'answer_status', 'limitations', 'human_validation']
    long_sheet.append(fields)
    for row in rows:
        response = _parse_response(row.get('response', ''))
        question = questions.get(row['question_id'], {})
        long_sheet.append([row['question_id'], question.get('question_text', ''), _country_code(row['country'], countries), row['country'], _summary(response), response.get('evidence_classification', ''), '; '.join(response.get('evidence_ids', [])), '; '.join(response.get('document_ids', [])), '; '.join(map(str, response.get('pages', []))), row.get('validation_status', 'candidate'), response.get('limitations', ''), ''])
    for index, country in enumerate(countries, start=2):
        sheet = workbook.create_sheet(f'{index:02d}_{country["country"]}'[:31])
        sheet.append(['question_id', 'pergunta', 'síntese da resposta', 'evidências', 'documentos/páginas', 'status'])
        for row in [item for item in rows if item['country'] == country['country']]:
            response = _parse_response(row.get('response', ''))
            sheet.append([row['question_id'], questions.get(row['question_id'], {}).get('question_text', ''), _summary(response), '; '.join(response.get('evidence_ids', [])), '; '.join(response.get('document_ids', [])), row.get('validation_status', 'candidate')])
    matrix = workbook.create_sheet('24_matriz_agregada')
    matrix.append(['question_id'] + [c['country'] for c in countries])
    for question_id in questions:
        values = {row['country']: _summary(_parse_response(row.get('response', ''))) for row in rows if row['question_id'] == question_id}
        matrix.append([question_id] + [values.get(c['country'], 'Não detectado') for c in countries])
    validation = workbook.create_sheet('25_controle_validacao')
    validation.append(['question_id', 'country', 'validation_status', 'evidence_ids'])
    validation.extend([[row['question_id'], row['country'], row.get('validation_status', 'candidate'), row.get('evidence_ids', '')] for row in rows])
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _parse_response(value: str) -> dict:
    try:
        parsed = json.loads(value)
        return parsed if isinstance(parsed, dict) else {}
    except (TypeError, json.JSONDecodeError):
        return {'answer_summary': value}


def _summary(response: dict) -> str:
    return response.get('answer_summary') or response.get('response') or response.get('summary') or ''


def _country_code(country: str, countries: list[dict]) -> str:
    return next((item['country_code'] for item in countries if item['country'] == country), '')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', default='analysis/countries')
    parser.add_argument('--out', default='02_perguntas_respostas.xlsx')
    args = parser.parse_args()
    print(build_workbook(args.input, args.out))