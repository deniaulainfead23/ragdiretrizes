from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from rag_project.corpus_registry import analysis_documents, load_registry
from rag_project.paths import ANALYSIS_DIR, INTERMEDIATE_DATA_DIR
from rag_project.validate_rag_answer import validate_rag_answer

DEFAULT_RUN_ID = "todos_paises_exploratorio_20260927_01"
DEFAULT_COUNTRY_NAMES = {
    "africa-do-sul": "África do Sul", "australia": "Austrália", "brasil": "Brasil",
    "canada": "Canadá", "chile": "Chile", "china": "China",
    "coreia-do-sul": "Coreia do Sul", "estonia": "Estônia", "eua": "Estados Unidos",
    "finlandia": "Finlândia", "gana": "Gana", "hong-kong": "Hong Kong",
    "irlanda": "Irlanda", "japao": "Japão", "nova-zelandia": "Nova Zelândia",
    "quenia": "Quênia", "reino-unido": "Reino Unido", "ruanda": "Ruanda",
    "singapura": "Singapura", "suica": "Suíça", "taiwan": "Taiwan", "uruguai": "Uruguai",
}


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _safe_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
    return value


def load_run_page_cache(run_id: str) -> tuple[dict[str, dict[str, str]], dict[str, Any]]:
    """Load extracted national page text to avoid reopening/OCRing PDFs per answer."""
    pipeline_run_dir = INTERMEDIATE_DATA_DIR / "pipeline_runs" / run_id
    dataset_path = pipeline_run_dir / "datasets" / "dataset_original_pages.jsonl"
    manifest_path = pipeline_run_dir / "run_manifest.json"
    if not dataset_path.is_file():
        raise FileNotFoundError(f"Dataset original da rodada não encontrado: {dataset_path}")

    page_cache: dict[str, dict[str, str]] = {}
    for entry, document in analysis_documents(load_registry(), include_pending_review=True):
        relative = (Path(str(entry.get("country", ""))) / str(document.get("file", ""))).as_posix()
        page_cache.setdefault(relative, {})

    records_loaded = 0
    with dataset_path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            record = json.loads(line)
            if record.get("source_scope") != "national":
                continue
            source_path = str(record.get("source_path", "")).replace("\\", "/")
            if source_path.startswith("corpus/"):
                source_path = source_path[len("corpus/"):]
            page = str(record.get("page_start", "")).strip()
            if source_path and page:
                page_cache.setdefault(source_path, {})[page] = str(record.get("source_text", ""))
                records_loaded += 1

    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.is_file() else {}
    return page_cache, {
        "page_text_source": str(dataset_path),
        "page_records_loaded": records_loaded,
        "page_documents_available": sum(bool(pages) for pages in page_cache.values()),
        "corpus_registry_sha256_at_run": manifest.get("registry_sha256", ""),
    }


def _write_workbook(path: Path, response_rows: list[dict], citation_rows: list[dict], summary: dict) -> None:
    workbook = Workbook()
    info = workbook.active
    info.title = "Instruções"
    for row in [
        ("Auditoria automática preliminar das respostas RAG", summary["run_id"]),
        ("Respostas processadas", summary["responses_processed"]),
        ("Status", str(summary["status_counts"])),
        ("Citações", str(summary["citation_status_counts"])),
        ("Modelo", summary["semantic_model"]),
        ("Limite", "Não é validação científica oficial. Toda resposta exige revisão humana, mesmo quando status=validado_preliminarmente."),
        ("Fonte textual", "A resposta pode ser paráfrase. A rotina exige passagem literal de apoio no PDF para confirmar uma paráfrase."),
        ("Curadoria", "O status do registro documental é preservado; esta auditoria não altera corpus_registry.yaml."),
    ]:
        info.append([_safe_cell(value) for value in row])

    response_headers = [
        "country", "question_id", "category_id", "question", "rag_status",
        "validation_status", "result_description", "semantic_decision",
        "semantic_reason", "citation_count", "verified_citation_count",
        "status_reasons", "human_review_required", "run_date",
    ]
    response_sheet = workbook.create_sheet("Respostas")
    response_sheet.append(response_headers)
    for row in response_rows:
        response_sheet.append([_safe_cell(row.get(field, "")) for field in response_headers])

    citation_headers = [
        "country", "question_id", "evidence_ref", "evidence_id", "citation_status",
        "document_id", "document_title", "page_cited", "physical_page",
        "source_path", "source_registry_status", "rag_excerpt",
        "verified_source_quote", "reason", "semantic_review_decision",
        "semantic_review_reason",
    ]
    citation_sheet = workbook.create_sheet("Citações")
    citation_sheet.append(citation_headers)
    for row in citation_rows:
        citation_sheet.append([_safe_cell(row.get(field, "")) for field in citation_headers])

    for sheet in (response_sheet, citation_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174A5B")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in sheet.columns:
            header = str(column[0].value or "").casefold()
            sheet.column_dimensions[column[0].column_letter].width = (
                55 if any(word in header for word in ("question", "excerpt", "quote", "reason", "description"))
                else min(max(len(str(column[0].value or "")) + 2, 14), 32)
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = str(sheet.cell(1, cell.column).value or "").casefold()
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=any(word in header for word in ("question", "excerpt", "quote", "reason", "description")),
                )

    info.column_dimensions["A"].width = 38
    info.column_dimensions["B"].width = 110
    for cell in info[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A5B")
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def validate_run(
    run_id: str = DEFAULT_RUN_ID,
    output_dir: Path | None = None,
    countries: set[str] | None = None,
    model: str = "gpt-4o-mini",
) -> dict[str, Any]:
    load_dotenv(Path(__file__).resolve().parent / ".env")
    run_dir = ANALYSIS_DIR / run_id / "countries"
    if not run_dir.is_dir():
        raise FileNotFoundError(f"Pasta de respostas da rodada não existe: {run_dir}")

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    destination = output_dir or (ANALYSIS_DIR / run_id / f"rag_validation_{timestamp}")
    destination.mkdir(parents=True, exist_ok=False)
    checkpoint_path = destination / "validation_results.jsonl"
    workbook_path = destination / "auditoria_respostas_rag.xlsx"

    planned: list[tuple[str, dict[str, str], list[dict[str, str]]]] = []
    for country_dir in sorted(run_dir.iterdir()):
        if not country_dir.is_dir() or (countries and country_dir.name not in countries):
            continue
        response_file = country_dir / "respostas.csv"
        evidence_file = country_dir / "evidencias.csv"
        if not response_file.is_file():
            continue
        responses = _read_csv(response_file)
        evidences = _read_csv(evidence_file) if evidence_file.is_file() else []
        for response in responses:
            qid = response.get("question_id", "")
            question_evidences = [
                evidence for evidence in evidences
                if evidence.get("question_id") == qid
                and evidence.get("response_id") == response.get("response_id")
            ]
            planned.append((country_dir.name, response, question_evidences))

    planned.sort(key=lambda item: (item[1].get("question_id", ""), item[0]))
    page_cache, page_cache_metadata = load_run_page_cache(run_id)
    response_results: list[dict] = []
    citation_results: list[dict] = []
    status_counts: Counter[str] = Counter()
    citation_counts: Counter[str] = Counter()

    with checkpoint_path.open("x", encoding="utf-8", newline="\n") as checkpoint:
        for index, (country, response, evidence_rows) in enumerate(planned, start=1):
            qid = response.get("question_id", "")
            payload = {
                "question_id": qid,
                "category_id": response.get("category_id", ""),
                "country": country,
                "response": response.get("response", ""),
                "evidences": evidence_rows,
                "validation_status": response.get("validation_status", ""),
            }
            try:
                result = validate_rag_answer(
                    response.get("question_text", ""),
                    country,
                    answer=json.dumps(payload, ensure_ascii=False),
                    question_id=qid,
                    category_id=response.get("category_id", ""),
                    model=model,
                    page_cache=page_cache,
                )
                result["rag_status"] = response.get("validation_status", "")
                result["run_date"] = response.get("run_date", "")
                result["validation_batch_stage"] = "automatic_preliminary_human_review_required"
                result["page_check_source"] = "run_original_dataset_pages"
            except Exception as exc:
                result = {
                    "run_id": run_id,
                    "country": country,
                    "question_id": qid,
                    "category_id": response.get("category_id", ""),
                    "question": response.get("question_text", ""),
                    "answer": response.get("response", ""),
                    "status": "conteudo_duvidoso",
                    "result_description": "Falha técnica durante a auditoria automática; revisão humana necessária.",
                    "validation_level": "automatic_preliminary_error",
                    "human_review_required": True,
                    "reasons": [f"Erro: {type(exc).__name__}: {str(exc)[:500]}"],
                    "status_reasons": [{"code": "validation_execution_error", "detail": str(exc)[:500]}],
                    "citation_checks": [],
                    "semantic_review": {"decision": "not_run", "reason": "Falha antes da revisão semântica."},
                    "rag_status": response.get("validation_status", ""),
                    "run_date": response.get("run_date", ""),
                    "validation_batch_stage": "automatic_preliminary_human_review_required",
                    "page_check_source": "run_original_dataset_pages",
                }

            response_results.append(result)
            status_counts[result.get("status", "conteudo_duvidoso")] += 1
            semantic = result.get("semantic_review") or {}
            for citation in result.get("citation_checks") or []:
                citation["question_id"] = qid
                citation["semantic_review_decision"] = semantic.get("decision", "")
                citation["semantic_review_reason"] = semantic.get("reason", "")
                citation_results.append(citation)
                citation_counts[citation.get("citation_status", "unresolved")] += 1
            checkpoint.write(json.dumps(result, ensure_ascii=False) + "\n")
            checkpoint.flush()
            if index % 10 == 0 or index == len(planned):
                print(f"Respostas auditadas: {index}/{len(planned)}", flush=True)

    _write_workbook(workbook_path, response_results, citation_results, {
        "run_id": run_id,
        "responses_processed": len(response_results),
        "status_counts": dict(status_counts),
        "citation_status_counts": dict(citation_counts),
        "semantic_model": model,
    })
    summary = {
        "run_id": run_id,
        "responses_processed": len(response_results),
        "citations_checked": len(citation_results),
        "status_counts": dict(status_counts),
        "citation_status_counts": dict(citation_counts),
        "page_cache_metadata": page_cache_metadata,
        "human_review_required_for_all": True,
        "checkpoint_path": str(checkpoint_path),
        "workbook_path": str(workbook_path),
    }
    (destination / "validation_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Audita respostas RAG existentes contra as páginas dos PDFs do corpus.")
    parser.add_argument("--run-id", default=DEFAULT_RUN_ID)
    parser.add_argument("--out", type=Path, default=None, help="Nova pasta de saída; não pode existir previamente")
    parser.add_argument("--countries", nargs="*", default=None, help="Opcional: slugs de países para validar apenas um subconjunto")
    parser.add_argument("--model", default="gpt-4o-mini")
    args = parser.parse_args()
    summary = validate_run(args.run_id, args.out, set(args.countries) if args.countries else None, args.model)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
