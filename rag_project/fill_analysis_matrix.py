from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "dados_derivados" / "analysis" / "official_run_20260912"
QUESTIONS_DIR = RUN / "review_exports" / "questions"
INPUT_XLSX = ROOT / "matriz_analise_rag_13x22_template.xlsx"
OUTPUT_XLSX = RUN / "matriz_analise_rag_13x22_preenchida.xlsx"

COUNTRY_MAP = {
    "africa-do-sul": "África do Sul",
    "australia": "Austrália",
    "brasil": "Brasil",
    "canada": "Canadá",
    "chile": "Chile",
    "china": "China",
    "coreia-do-sul": "Coreia do Sul",
    "estonia": "Estônia",
    "eua": "EUA",
    "finlandia": "Finlândia",
    "gana": "Gana",
    "hong-kong": "Hong Kong",
    "irlanda": "Irlanda",
    "japao": "Japão",
    "nova-zelandia": "Nova Zelândia",
    "quenia": "Quênia",
    "reino-unido": "Reino Unido",
    "ruanda": "Ruanda",
    "singapura": "Singapura",
    "suica": "Suíça",
    "taiwan": "Taiwan",
    "uruguai": "Uruguai",
}

if not INPUT_XLSX.exists():
    raise FileNotFoundError(
        f"Planilha modelo não encontrada: {INPUT_XLSX}\n"
        "Coloque matriz_analise_rag_13x22_template.xlsx na raiz do repositório."
    )

wb = load_workbook(INPUT_XLSX)
matrix = wb["Matriz_13x22"]
status_sheet = wb["Status"]
evidence_sheet = wb["Evidências"]

country_columns = {}
for col in range(3, matrix.max_column + 1):
    name = matrix.cell(row=1, column=col).value
    if name:
        country_columns[str(name).strip()] = col

for qnum in range(1, 14):
    qid = f"Q{qnum:02d}"
    file = QUESTIONS_DIR / f"{qid}.csv"
    if not file.exists():
        print(f"{qid}: arquivo não encontrado: {file}")
        continue

    df = pd.read_csv(file)
    row = qnum + 1

    for _, rec in df.iterrows():
        country_key = str(rec.get("country", "")).strip()
        country_name = COUNTRY_MAP.get(country_key)
        if not country_name:
            print(f"{qid}: país desconhecido: {country_key}")
            continue

        col = country_columns.get(country_name)
        if not col:
            print(f"{qid}: coluna não encontrada para {country_name}")
            continue

        response = rec.get("response", "")
        status = rec.get("validation_status", "")

        matrix.cell(row=row, column=col).value = "" if pd.isna(response) else str(response)
        status_sheet.cell(row=row, column=col).value = "" if pd.isna(status) else str(status)

    print(f"{qid}: preenchida")

# A matriz de respostas e a aba de evidências têm funções diferentes:
# respostas permanecem candidatas, enquanto cada trecho recuperado fica rastreável.
evidence_sheet.delete_rows(2, evidence_sheet.max_row)
evidence_columns = {
    str(evidence_sheet.cell(row=1, column=column).value).strip(): column
    for column in range(1, evidence_sheet.max_column + 1)
    if evidence_sheet.cell(row=1, column=column).value
}
evidence_rows = []
for country_key, country_name in COUNTRY_MAP.items():
    evidence_file = RUN / country_key / "evidencias.csv"
    if not evidence_file.exists():
        print(f"Evidências não encontradas: {evidence_file}")
        continue
    evidence_df = pd.read_csv(evidence_file).fillna("")
    evidence_df = evidence_df[evidence_df["question_id"].astype(str).str.match(r"^Q(?:0[1-9]|1[0-3])$")]
    for _, record in evidence_df.iterrows():
        evidence_status = str(record.get("validation_status", "")).strip() or "candidate"
        evidence_rows.append({
            "question_id": record.get("question_id", ""),
            "country": country_name,
            "evidence_id": record.get("evidence_id", ""),
            "document_id": record.get("document_id", ""),
            "document_title": record.get("document_title", ""),
            "page_start": record.get("page_start", ""),
            "page_end": record.get("page_end", ""),
            "source_text": record.get("source_text", ""),
            "validation_status": evidence_status,
            "validator_notes": record.get("review_notes", ""),
        })

for row_number, record in enumerate(evidence_rows, start=2):
    for field, column in evidence_columns.items():
        evidence_sheet.cell(row=row_number, column=column).value = record.get(field, "")

print(f"Evidências copiadas para a matriz: {len(evidence_rows)}")

OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT_XLSX)

print("\nPlanilha criada:")
print(OUTPUT_XLSX)
