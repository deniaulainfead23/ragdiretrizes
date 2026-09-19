from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

RUN = ROOT / "analysis" / "official_run_20260912"
QUESTIONS_DIR = RUN / "review_exports" / "questions"

INPUT_XLSX = ROOT / "matriz_analise_rag_13x22_template.xlsx"
OUTPUT_XLSX = ROOT / "analysis" / "official_run_20260912" / "matriz_analise_rag_13x22_preenchida.xlsx"

COUNTRY_MAP = {
    "africa-do-sul": "África do Sul",
    "australia": "Austrália",
    "brasil": "Brasil",
    "canada": "Canadá",
    "chile": "Chile",
    "china": "China",
    "coreia-do-sul": "Coreia do Sul",
    "estonia": "Estônia",
    "eua": "EUA",
    "finlandia": "Finlândia",
    "gana": "Gana",
    "hong-kong": "Hong Kong",
    "irlanda": "Irlanda",
    "japao": "Japão",
    "nova-zelandia": "Nova Zelândia",
    "quenia": "Quênia",
    "reino-unido": "Reino Unido",
    "ruanda": "Ruanda",
    "singapura": "Singapura",
    "suica": "Suíça",
    "taiwan": "Taiwan",
    "uruguai": "Uruguai",
}

wb = load_workbook(INPUT_XLSX)

matrix = wb["Matriz_13x22"]
status_sheet = wb["Status"]

# Descobre a coluna de cada país pelo cabeçalho da planilha
country_columns = {}

for col in range(3, matrix.max_column + 1):
    name = matrix.cell(row=1, column=col).value
    if name:
        country_columns[name] = col

for qnum in range(1, 14):

    qid = f"Q{qnum:02d}"
    file = QUESTIONS_DIR / f"{qid}.csv"

    if not file.exists():
        print(f"{qid}: arquivo não encontrado")
        continue

    df = pd.read_csv(file)

    # linha correspondente na planilha:
    # Q01 = linha 2, Q02 = linha 3...
    row = qnum + 1

    for _, rec in df.iterrows():

        country_key = str(rec.get("country", "")).strip()
        country_name = COUNTRY_MAP.get(country_key)

        if not country_name:
            print(f"{qid}: país desconhecido: {country_key}")
            continue

        col = country_columns.get(country_name)

        if not col:
            print(f"{qid}: coluna não encontrada para {country_name}")
            continue

        response = rec.get("response", "")
        status = rec.get("validation_status", "")

        if pd.isna(response):
            response = ""

        if pd.isna(status):
            status = ""

        matrix.cell(row=row, column=col).value = str(response)
        status_sheet.cell(row=row, column=col).value = str(status)

    print(f"{qid}: preenchida")

OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)

wb.save(OUTPUT_XLSX)

print()
print("Planilha criada:")
print(OUTPUT_XLSX)